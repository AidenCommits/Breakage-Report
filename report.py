from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from math import ceil, floor, log10

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter, MultipleLocator

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from analytics import calculate_analytics, get_financial_events
from config import LOCAL_TIMEZONE, REPORT_DIR
from database import get_connection

NAVY = "17365D"
LIGHT_BLUE = "EAF2F8"
WHITE = "FFFFFF"
BORDER_COLOR = "9DB1C5"
CURRENCY_FORMAT = '"$"#,##0.00;[Red]("$"#,##0.00)'


def configure_printing(ws, repeat_headers=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.45, bottom=0.45,
        header=0.2, footer=0.2,
    )
    ws.oddFooter.center.text = "Page &P of &N"
    if repeat_headers:
        ws.print_title_rows = "1:1"


def style_table(ws, header_row=1):
    configure_printing(ws)
    thin = Side(style="thin", color=BORDER_COLOR)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == header_row:
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
            else:
                cell.fill = PatternFill(
                    "solid", fgColor=WHITE if cell.row % 2 == 0 else LIGHT_BLUE
                )
                cell.font = Font(name="Aptos", size=10)
    ws.row_dimensions[header_row].height = 30
    for row in range(header_row + 1, ws.max_row + 1):
        ws.row_dimensions[row].height = 26
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        longest = max(
            (len(str(ws.cell(row, col).value or "")) for row in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(longest + 3, 16), 55)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = f"A{header_row + 1}"
    ws.print_area = ws.dimensions


def get_damage_register():
    with get_connection() as connection:
        rows = connection.execute("""
            SELECT job_name, description, room, quantity,
                   cost_type, current_cost, status
            FROM damaged_items
            ORDER BY job_name, description
        """).fetchall()
    return [dict(row) for row in rows]


def build_executive_summary(wb, results):
    ws = wb.active
    ws.title = "Executive Summary"
    start = results["week_start"].strftime("%b %d, %Y")
    end = (results["week_end"] - timedelta(days=1)).strftime("%b %d, %Y")
    metrics = [
        ("Reporting Period", f"{start} to {end}"),
        ("Previous Week Losses", results["weekly_total"]),
        ("Cumulative Losses", results["cumulative_total"]),
        ("Average Weekly Losses", results["average_weekly_loss"]),
        ("Annualized Projection (Previous Week x 52)", results["annualized_projection"]),
        ("Annualized Historical", results["annualized_historical"]),
    ]
    ws.append(["Financial Metric", "Value"])
    for label, value in metrics:
        ws.append([label, float(value) if isinstance(value, Decimal) else value])
    style_table(ws)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 38
    for row in range(3, ws.max_row + 1):
        ws.cell(row, 2).number_format = CURRENCY_FORMAT
    ws["B3"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
    ws["B6"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)


def build_job_breakdown(wb, results):
    ws = wb.create_sheet("Job Breakdown")
    ws.append(["Job", "Previous Week Loss"])
    for job, amount in sorted(
        results["weekly_by_job"].items(), key=lambda item: item[1], reverse=True
    ):
        ws.append([job, float(amount)])
    style_table(ws)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 26
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).number_format = CURRENCY_FORMAT


def build_damage_register(wb, records):
    ws = wb.create_sheet("Damage Register")
    ws.append(["Job", "Item", "Room", "Quantity", "Cost Type", "Cost", "Status"])
    for record in records:
        raw_cost = record["current_cost"]
        cost = float(Decimal(str(raw_cost))) if raw_cost not in (None, "") else None
        ws.append([
            record["job_name"], record["description"], record["room"],
            record["quantity"], record["cost_type"], cost, record["status"],
        ])
    style_table(ws)
    for column, width in {"A": 30, "B": 40, "C": 24, "D": 14,
                          "E": 20, "F": 18, "G": 20}.items():
        ws.column_dimensions[column].width = width
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 6).number_format = CURRENCY_FORMAT


def build_historical_trends(wb, results):
    ws = wb.create_sheet("Historical Trends")
    ws.append(["Week Starting", "Weekly Loss", "Cumulative Loss"])
    running_total = Decimal("0")
    for week in results["weekly_history"]:
        running_total += week["total"]
        ws.append([
            week["week_start"].strftime("%b %d, %Y"),
            float(week["total"]), float(running_total),
        ])
    style_table(ws)
    for column in (2, 3):
        for row in range(2, ws.max_row + 1):
            ws.cell(row, column).number_format = CURRENCY_FORMAT
    for column in "ABC":
        ws.column_dimensions[column].width = 26


def choose_axis_scale(values):
    """Return a readable Y-axis ceiling and tick spacing."""
    peak = max([0.0] + [float(value) for value in values])
    if peak <= 0:
        return 100.0, 25.0
    target = peak * 1.12 / 4
    magnitude = 10 ** floor(log10(target))
    step = next(
        (factor * magnitude for factor in (1, 2, 2.5, 5, 10)
         if factor * magnitude >= target),
        10 * magnitude,
    )
    ceiling = max(step, ceil(peak * 1.12 / step) * step)
    return float(ceiling), float(step)


def make_chart_image(weeks, values, title, chart_kind):
    """Create a print-quality PNG in memory for embedding in Excel."""
    fig, ax = plt.subplots(figsize=(12.5, 6.4), dpi=160)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    positions = list(range(len(weeks)))
    ceiling, step = choose_axis_scale(values)
    ax.set_ylim(0, ceiling)
    ax.yaxis.set_major_locator(MultipleLocator(step))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda number, _: f"${number:,.0f}"))
    ax.grid(axis="y", color="#DCE4EC", linewidth=1)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#9DAAB8")
    ax.spines["bottom"].set_color("#9DAAB8")
    ax.tick_params(axis="both", labelsize=12, length=0, pad=10)
    ax.set_title(title, fontsize=20, fontweight="bold", color="#17365D", pad=25)
    ax.set_ylabel("Losses ($)", fontsize=13, labelpad=12)
    ax.set_xlabel("Week Starting", fontsize=13, labelpad=14)
    ax.set_xticks(positions)
    ax.set_xticklabels(weeks, rotation=35 if len(weeks) > 8 else 0,
                       ha="right" if len(weeks) > 8 else "center")
    ax.set_xlim(-2, 2) if len(weeks) == 1 else ax.set_xlim(-0.7, len(weeks) - 0.3)

    if chart_kind == "weekly":
        ax.bar(positions, values, width=0.55, color="#477FB8")
    else:
        if len(positions) > 1:
            ax.plot(positions, values, color="#4F9A65", linewidth=3,
                    marker="o", markersize=9)
        else:
            ax.scatter(positions, values, color="#4F9A65", s=110, zorder=4)

    for x, value in zip(positions, values):
        if value != 0:
            ax.annotate(f"${value:,.2f}", (x, value), xytext=(0, 10),
                        textcoords="offset points", ha="center", va="bottom",
                        fontsize=11, fontweight="bold", color="#263746")

    fig.subplots_adjust(left=0.11, right=0.97, top=0.85, bottom=0.19)
    buffer = BytesIO()
    fig.savefig(buffer, format="png", dpi=160, facecolor="white")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def configure_chart_sheet(ws):
    configure_printing(ws, repeat_headers=False)
    ws.page_setup.fitToHeight = 1
    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = 10
    for row in range(1, 36):
        ws.row_dimensions[row].height = 20
    ws.print_area = "A1:P35"


def show_no_data_message(ws):
    ws.merge_cells("B10:O13")
    cell = ws["B10"]
    cell.value = "No financial losses recorded for this reporting period."
    cell.font = Font(name="Aptos", size=18, bold=True, color="666666")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def insert_chart_image(ws, weeks, values, title, kind, image_buffers):
    image_buffer = make_chart_image(weeks, values, title, kind)
    image_buffers.append(image_buffer)  # Keep the PNG available until workbook save.
    chart_image = ExcelImage(image_buffer)
    chart_image.width = 950
    chart_image.height = 486
    ws.add_image(chart_image, "B2")


def build_weekly_loss_chart(wb, results, image_buffers):
    ws = wb.create_sheet("Weekly Loss Chart")
    configure_chart_sheet(ws)
    history = results["weekly_history"]
    if not history:
        show_no_data_message(ws)
        return
    weeks = [entry["week_start"].strftime("%b %d, %Y") for entry in history]
    values = [float(entry["total"]) for entry in history]
    insert_chart_image(ws, weeks, values, "Weekly Financial Losses", "weekly", image_buffers)


def build_cumulative_loss_chart(wb, results, image_buffers):
    ws = wb.create_sheet("Cumulative Loss Chart")
    configure_chart_sheet(ws)
    history = results["weekly_history"]
    if not history:
        show_no_data_message(ws)
        return
    weeks = [entry["week_start"].strftime("%b %d, %Y") for entry in history]
    running = 0.0
    values = []
    for entry in history:
        running += float(entry["total"])
        values.append(running)
    insert_chart_image(ws, weeks, values, "Cumulative Financial Losses", "cumulative", image_buffers)


def generate_report(now=None):
    """Generate the complete Excel management report."""
    events = get_financial_events()
    results = calculate_analytics(events, now=now)
    damage_records = get_damage_register()
    
    # results = calculate_analytics(
    #     events,
    #     now=datetime(
    #         2026, 10, 12, 9, 0,
    #         tzinfo=LOCAL_TIMEZONE
    #     )
    # )

    print("\n--- REPORT SUMMARY ---")
    print("Reporting period:", results["week_start"], "to", results["week_end"])
    print("Weekly total:", results["weekly_total"])
    print("Annualized projection:", results["annualized_projection"])
    print("--- END SUMMARY ---\n")

    wb = Workbook()
    image_buffers = []
    build_executive_summary(wb, results)
    build_job_breakdown(wb, results)
    build_damage_register(wb, damage_records)
    build_historical_trends(wb, results)
    build_weekly_loss_chart(wb, results, image_buffers)
    build_cumulative_loss_chart(wb, results, image_buffers)

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(LOCAL_TIMEZONE).strftime("%Y-%m-%d_%H%M%S")
    output_path = REPORT_DIR / f"Breakage_Report_{timestamp}.xlsx"
    try:
        wb.save(output_path)
    finally:
        for buffer in image_buffers:
            buffer.close()
    print(f"Excel report generated: {output_path}")
    return output_path


if __name__ == "__main__":
    generate_report()
